import pytest
import uuid
import openpyxl
from io import BytesIO
from unittest.mock import AsyncMock, MagicMock, patch
from httpx import AsyncClient, ASGITransport
from sqlalchemy import select

from app.main import app
from app.db.session import AsyncSessionLocal
from app.models.project import Project
from app.models.membership import UserTenantMembership
from app.schemas.dashboard import Widget
from app.core.llm_provider import LLMProvider, LLMResponse, UsageInfo
from app.services.export_service import ExportService

class MockLLMProvider(LLMProvider):
    async def complete(
        self,
        messages,
        tenant_id,
        operation,
        response_format=None,
        temperature=0.7
    ):
        return LLMResponse(
            content='{"section_names": ["KPI Summary", "Campaign Analytics"], "sheet_order": ["Summary", "Details"]}',
            usage=UsageInfo(prompt_tokens=10, completion_tokens=15, total_tokens=25),
            model="gpt-4o-mini"
        )

class MockFailingLLMProvider(LLMProvider):
    async def complete(
        self,
        messages,
        tenant_id,
        operation,
        response_format=None,
        temperature=0.7
    ):
        raise Exception("LLM Provider Timeout or Outage")

@pytest.fixture
async def auth_client():
    email = f"user_{uuid.uuid4().hex[:8]}@example.com"
    password = "secretpassword"
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        # Register user
        register_res = await client.post(
            "/api/v1/auth/register",
            json={"email": email, "password": password, "full_name": "Test User"}
        )
        user_id = register_res.json()["id"]
        
        # Login
        login_res = await client.post(
            "/api/v1/auth/login",
            data={"username": email, "password": password}
        )
        token = login_res.json()["access_token"]
        headers = {"Authorization": f"Bearer {token}"}
        
        # Retrieve tenant_id directly from DB membership
        async with AsyncSessionLocal() as session:
            res = await session.execute(
                select(UserTenantMembership).where(UserTenantMembership.user_id == uuid.UUID(user_id))
            )
            membership = res.scalars().first()
            tenant_id = membership.tenant_id
            
        yield client, headers, tenant_id, uuid.UUID(user_id)

@pytest.fixture
def sample_widgets():
    return [
        Widget(
            i="w1",
            x=0, y=0, w=6, h=4,
            chart="line",
            dataKey="ctr",
            title="CTR Performance Chart",
            data=[
                {"name": "2026-05-10", "ctr": 2.5, "spend": 100},
                {"name": "2026-05-11", "ctr": 3.1, "spend": 120}
            ]
        ),
        Widget(
            i="w2",
            x=6, y=0, w=6, h=4,
            chart="number",
            dataKey="spend",
            title="Total Spend KPI",
            data=[
                {"name": "2026-05-10", "ctr": 2.5, "spend": 100},
                {"name": "2026-05-11", "ctr": 3.1, "spend": 120}
            ]
        )
    ]

@pytest.mark.asyncio
async def test_export_generates_xlsx_bytes(sample_widgets):
    llm = MockLLMProvider()
    db_mock = AsyncMock()
    
    result = await ExportService.generate_report(
        widgets=sample_widgets,
        project_name="My Demo Project",
        tenant_id="test-tenant-id",
        llm=llm,
        db=db_mock
    )
    
    assert isinstance(result, bytes)
    assert len(result) > 0
    
    # Try opening the workbook using openpyxl
    wb = openpyxl.load_workbook(BytesIO(result))
    assert len(wb.sheetnames) >= 1
    
    # Assert there is actual sheet data
    ws = wb.active
    assert ws.max_row >= 1

@pytest.mark.asyncio
async def test_export_sheet_contains_widget_titles(sample_widgets):
    llm = MockLLMProvider()
    db_mock = AsyncMock()
    
    result = await ExportService.generate_report(
        widgets=sample_widgets,
        project_name="Campaign Brand Boost",
        tenant_id="test-tenant-id",
        llm=llm,
        db=db_mock
    )
    
    wb = openpyxl.load_workbook(BytesIO(result))
    
    # Let's search all cells in all sheets for the widget titles
    found_title = False
    for name in wb.sheetnames:
        ws = wb[name]
        for row in ws.iter_rows(values_only=True):
            for val in row:
                if val and ("CTR Performance Chart" in str(val) or "Total Spend KPI" in str(val)):
                    found_title = True
                    break
    assert found_title, "Widget titles should be present in the generated sheets"

@pytest.mark.asyncio
async def test_export_works_without_llm(sample_widgets):
    llm = MockFailingLLMProvider()
    db_mock = AsyncMock()
    
    # This should not raise any exceptions even though LLM fails
    result = await ExportService.generate_report(
        widgets=sample_widgets,
        project_name="Campaign Brand Boost",
        tenant_id="test-tenant-id",
        llm=llm,
        db=db_mock
    )
    
    assert isinstance(result, bytes)
    assert len(result) > 0
    
    wb = openpyxl.load_workbook(BytesIO(result))
    assert len(wb.sheetnames) >= 1

@pytest.mark.asyncio
async def test_export_endpoint_returns_file(auth_client, sample_widgets):
    client, headers, tenant_id, _ = auth_client
    
    # Create project using the existing DB session Local
    async with AsyncSessionLocal() as session:
        project = Project(
            id=uuid.uuid4(),
            tenant_id=tenant_id,
            name="Export Verification Project"
        )
        session.add(project)
        await session.commit()
        project_id = project.id

    # Mock the ExportService.generate_report method
    with patch("app.services.export_service.ExportService.generate_report") as mock_gen:
        mock_gen.return_value = b"mock_xlsx_bytes_data"
        
        response = await client.get(
            f"/api/v1/projects/{project_id}/dashboard/export",
            headers=headers
        )
        
        assert response.status_code == 200
        assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert "attachment" in response.headers["content-disposition"]
        assert f"filename={project_id}_report.xlsx" in response.headers["content-disposition"]
        assert response.content == b"mock_xlsx_bytes_data"
        
        # Verify the service was called correctly
        mock_gen.assert_called_once()


@pytest.mark.asyncio
async def test_export_sanitizes_invalid_sheet_names():
    class CustomMockLLMProvider(LLMProvider):
        async def complete(self, messages, tenant_id, operation, response_format=None, temperature=0.7):
            return LLMResponse(
                content='{"section_names": ["S1"], "sheet_order": ["Q1:Sales/ROI*2025", "Data[2025]"]}',
                usage=UsageInfo(prompt_tokens=10, completion_tokens=15, total_tokens=25),
                model="gpt-4o-mini"
            )
            
    llm = CustomMockLLMProvider()
    db_mock = AsyncMock()
    widgets = []
    
    result = await ExportService.generate_report(
        widgets=widgets,
        project_name="Demo",
        tenant_id="test",
        llm=llm,
        db=db_mock
    )
    
    assert isinstance(result, bytes)
    assert len(result) > 0
    
    wb = openpyxl.load_workbook(BytesIO(result))
    assert len(wb.sheetnames) >= 1
    
    invalid_chars = set("[]:*?/\\")
    for sheet_name in wb.sheetnames:
        assert not any(c in invalid_chars for c in sheet_name)


@pytest.mark.asyncio
async def test_export_truncates_oversized_data():
    llm = MockLLMProvider()
    db_mock = AsyncMock()
    
    huge_data = [{"col1": i, "col2": f"val_{i}"} for i in range(10_000)]
    widget = Widget(
        i="huge_w",
        x=0, y=0, w=6, h=4,
        chart="line",
        dataKey="col1",
        title="Huge Widget",
        data=huge_data
    )
    
    result = await ExportService.generate_report(
        widgets=[widget],
        project_name="Huge Project",
        tenant_id="test",
        llm=llm,
        db=db_mock
    )
    
    assert isinstance(result, bytes)
    
    wb = openpyxl.load_workbook(BytesIO(result))
    ws = wb.active or wb.worksheets[0]
    
    from app.services.export_service import MAX_ROWS_PER_WIDGET
    assert ws.max_row <= MAX_ROWS_PER_WIDGET + 10


@pytest.mark.asyncio
async def test_export_endpoint_response_opens_as_xlsx(auth_client):
    client, headers, tenant_id, _ = auth_client
    
    async with AsyncSessionLocal() as session:
        project = Project(
            id=uuid.uuid4(),
            tenant_id=tenant_id,
            name="Real Export Verification Project"
        )
        session.add(project)
        await session.commit()
        project_id = project.id

    from app.api.v1.deps import get_llm_provider
    app.dependency_overrides[get_llm_provider] = lambda: MockLLMProvider()
    try:
        response = await client.get(
            f"/api/v1/projects/{project_id}/dashboard/export",
            headers=headers
        )
        
        assert response.status_code == 200
        assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        
        wb = openpyxl.load_workbook(BytesIO(response.content))
        assert len(wb.worksheets) >= 1
    finally:
        app.dependency_overrides.pop(get_llm_provider, None)


@pytest.mark.asyncio
async def test_export_empty_widgets_list():
    llm = MockLLMProvider()
    db_mock = AsyncMock()
    
    result = await ExportService.generate_report(
        widgets=[],
        project_name="Empty Project",
        tenant_id="test",
        llm=llm,
        db=db_mock
    )
    
    assert isinstance(result, bytes)
    wb = openpyxl.load_workbook(BytesIO(result))
    assert len(wb.sheetnames) >= 1


@pytest.mark.asyncio
async def test_export_malformed_widget_data():
    llm = MockLLMProvider()
    db_mock = AsyncMock()
    
    malformed_data = [{"col": None}, {"col": {"nested": "dict"}}, {}]
    widget = Widget(
        i="w_malformed",
        x=0, y=0, w=6, h=4,
        chart="line",
        dataKey="col",
        title="Malformed Widget",
        data=malformed_data
    )
    
    result = await ExportService.generate_report(
        widgets=[widget],
        project_name="Malformed Project",
        tenant_id="test",
        llm=llm,
        db=db_mock
    )
    
    assert isinstance(result, bytes)
    wb = openpyxl.load_workbook(BytesIO(result))
    assert len(wb.sheetnames) >= 1


@pytest.mark.asyncio
async def test_export_hard_stop_after_max_total_cells():
    llm = MockLLMProvider()
    db_mock = AsyncMock()
    
    from app.services.export_service import MAX_TOTAL_CELLS
    num_cols = 25
    num_rows = 5000
    huge_data = [{f"col{j}": f"val_{i}_{j}" for j in range(num_cols)} for i in range(num_rows)]
    
    w1 = Widget(
        i="w1", x=0, y=0, w=6, h=4, chart="line", dataKey="col0",
        title="Widget One - Large Data",
        data=huge_data
    )
    w2 = Widget(
        i="w2", x=6, y=0, w=6, h=4, chart="bar", dataKey="val",
        title="SHOULD_NOT_BE_PROCESSED_WIDGET_TWO",
        data=[{"val": 100}]
    )
    
    result = await ExportService.generate_report(
        widgets=[w1, w2],
        project_name="Hard Stop Project",
        tenant_id="test",
        llm=llm,
        db=db_mock
    )
    
    wb = openpyxl.load_workbook(BytesIO(result))
    found_w2 = False
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for cell_val in row:
                if cell_val and "SHOULD_NOT_BE_PROCESSED_WIDGET_TWO" in str(cell_val):
                    found_w2 = True
    assert not found_w2, "Widget 2 should not be processed after MAX_TOTAL_CELLS was reached"


@pytest.mark.asyncio
async def test_export_duplicate_sanitized_sheet_names():
    class DuplicateSheetLLMProvider(LLMProvider):
        async def complete(self, messages, tenant_id, operation, response_format=None, temperature=0.7):
            return LLMResponse(
                content='{"section_names": ["S1"], "sheet_order": ["A/B", "A:B", "A?B", "A/B"]}',
                usage=UsageInfo(prompt_tokens=10, completion_tokens=15, total_tokens=25),
                model="gpt-4o-mini"
            )
            
    llm = DuplicateSheetLLMProvider()
    db_mock = AsyncMock()
    
    result = await ExportService.generate_report(
        widgets=[],
        project_name="Dedupe Sheet Names Project",
        tenant_id="test",
        llm=llm,
        db=db_mock
    )
    
    wb = openpyxl.load_workbook(BytesIO(result))
    assert "A_B" in wb.sheetnames
    assert "A_B_2" in wb.sheetnames
    assert "A_B_3" in wb.sheetnames
    assert "A_B_4" in wb.sheetnames


@pytest.mark.asyncio
async def test_export_max_widgets_truncation():
    llm = MockLLMProvider()
    db_mock = AsyncMock()
    
    from app.services.export_service import MAX_WIDGETS
    widgets = []
    for i in range(MAX_WIDGETS + 10):
        widgets.append(Widget(
            i=f"w_{i}", x=0, y=0, w=1, h=1, chart="number", dataKey="val",
            title=f"Widget Title {i}",
            data=[]
        ))
        
    result = await ExportService.generate_report(
        widgets=widgets,
        project_name="Max Widgets Truncation",
        tenant_id="test",
        llm=llm,
        db=db_mock
    )
    
    wb = openpyxl.load_workbook(BytesIO(result))
    found_truncated_widget = False
    target_truncated_title = f"Widget Title {MAX_WIDGETS + 2}"
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for cell_val in row:
                if cell_val and target_truncated_title in str(cell_val):
                    found_truncated_widget = True
    assert not found_truncated_widget, "Widgets past MAX_WIDGETS should be truncated"

